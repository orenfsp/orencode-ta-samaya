from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from math import ceil
from typing import Any, Literal
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

INK, MUTED, PRIMARY = "1B1C21", "5A5D6B", "3D2BFF"
HEAD = PatternFill("solid", fgColor="E8E6FF")
ZEBRA = PatternFill("solid", fgColor="F6F6FA")
LINE = Side(style="thin", color="D5D7E0")
GRID = Border(left=LINE, right=LINE, top=LINE, bottom=LINE)
ORG = "ФСП Оренбургской области"

Kind = Literal["text", "int", "date", "percent", "center", "wrap"]


@dataclass(frozen=True, slots=True)
class Column:
    title: str
    kind: Kind = "text"
    width: int | None = None


def book() -> Workbook:
    """Пустая книга без листа по умолчанию: листы добавляет sheet()."""
    result = Workbook()
    if result.active is not None:
        result.remove(result.active)
    return result


def sheet(workbook: Workbook, title: str) -> Worksheet:
    """Новый лист с допустимым именем и настройкой печати: A4 альбомная, по ширине страницы."""
    clean = "".join(char for char in title if char not in "[]:*?/\\")[:31]
    result = workbook.create_sheet(clean)
    result.sheet_view.showGridLines = False
    result.page_setup.orientation = "landscape"
    result.page_setup.paperSize = result.PAPERSIZE_A4
    result.page_setup.fitToWidth = 1
    result.page_setup.fitToHeight = 0
    result.sheet_properties.pageSetUpPr.fitToPage = True
    result.print_options.horizontalCentered = True
    return result


def header(target: Worksheet, title: str, facts: Sequence[tuple[str, Any]] = ()) -> int:
    """Шапка листа: организация, заголовок, строки «подпись: значение» во всю ширину, дата формирования. Возвращает первую свободную строку."""
    target.cell(1, 1, ORG).font = Font(size=9, color=MUTED)
    target.cell(2, 1, title).font = Font(bold=True, size=15, color=INK)
    row = 3
    stamp = datetime.now(ZoneInfo("Asia/Yekaterinburg")).strftime("%d.%m.%Y %H:%M")
    for label, value in [*facts, ("Сформирован", stamp)]:
        text = value.strftime("%d.%m.%Y") if isinstance(value, date) else _value(value)
        target.cell(row, 1, f"{label}: {text}").font = Font(color=MUTED)
        row += 1
    return row + 1


def caption(target: Worksheet, row: int, text: str) -> int:
    """Подзаголовок раздела на листе; возвращает строку под ним."""
    target.cell(row, 1, text).font = Font(bold=True, size=12, color=INK)
    return row + 1


def table(
    target: Worksheet,
    row: int,
    columns: Sequence[Column],
    rows: Sequence[Sequence[Any]],
    *,
    total: Sequence[Any] | None = None,
    freeze: bool = True,
    empty: str = "Нет данных",
) -> int:
    """Таблица с фирменной шапкой, зеброй и сеткой; главная таблица листа (freeze) — с автофильтром и закреплённой шапкой; empty="" — без строки «нет данных». Возвращает строку после таблицы."""
    if freeze:
        available = sum(column.width or 20 for column in columns)
        for header_row in range(1, row):
            cell = target.cell(header_row, 1)
            if cell.value is not None and len(columns) > 1 and not any(target.cell(header_row, index).value is not None for index in range(2, len(columns) + 1)):
                target.merge_cells(start_row=header_row, start_column=1, end_row=header_row, end_column=len(columns))
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                target.row_dimensions[header_row].height = max(20, ceil(len(str(cell.value)) / max(available * .85, 1)) * (22 if header_row == 2 else 16))
    for index, column in enumerate(columns, start=1):
        cell = target.cell(row, index, column.title)
        cell.font = Font(bold=True, color=INK)
        cell.fill = HEAD
        cell.border = GRID
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    target.row_dimensions[row].height = 30
    head = row
    widths = [max(len(column.title) * 0.9, 8) for column in columns]
    if not rows and empty:
        target.cell(row + 1, 1, empty).font = Font(italic=True, color=MUTED)
        row += 1
    for offset, values in enumerate(rows):
        row += 1
        for index, (column, value) in enumerate(zip(columns, values, strict=False), start=1):
            cell = target.cell(row, index, _value(value))
            _format(cell, column.kind, value)
            cell.border = GRID
            if offset % 2:
                cell.fill = ZEBRA
            widths[index - 1] = max(widths[index - 1], min(len(_text(value)), 60))
    if total is not None:
        row += 1
        for index, (column, value) in enumerate(zip(columns, total, strict=False), start=1):
            cell = target.cell(row, index, _value(value))
            _format(cell, column.kind, value)
            cell.font = Font(bold=True, color=INK)
            cell.border = Border(top=Side(style="medium", color=INK), bottom=LINE, left=LINE, right=LINE)
    if freeze:
        target.auto_filter.ref = f"A{head}:{get_column_letter(len(columns))}{max(head + len(rows), head + 1)}"
    if freeze:
        target.freeze_panes = f"A{head + 1}"
        target.print_title_rows = f"{head}:{head}"
    for index, column in enumerate(columns, start=1):
        letter = get_column_letter(index)
        width = column.width or round(widths[index - 1] + 3)
        target.column_dimensions[letter].width = max(width, target.column_dimensions[letter].width or 0)
    return row + 2


def link(target: Worksheet, row: int, column: int, url: str, text: str) -> None:
    """Ссылка в ячейке: кликается в Excel и LibreOffice."""
    cell = target.cell(row, column, text)
    cell.hyperlink = url
    cell.font = Font(color=PRIMARY, underline="single")


def signatures(target: Worksheet, row: int, people: Sequence[tuple[str, str]]) -> int:
    """Строки подписей «должность ______ ФИО» под таблицей."""
    for position, name in people:
        target.cell(row, 1, position).font = Font(color=INK)
        line = target.cell(row, 3, "")
        line.border = Border(bottom=Side(style="thin", color=INK))
        target.cell(row, 4, name).font = Font(color=INK)
        row += 2
    return row


def save(workbook: Workbook) -> bytes:
    """Книга в байты для ответа API."""
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _value(value: Any) -> Any:
    """Значение для ячейки: None — пусто, даты и числа как есть, остальное строкой."""
    if value is None:
        return ""
    if isinstance(value, (int, float, date)):
        return value
    return str(value)


def _text(value: Any) -> str:
    """Текст значения для подбора ширины колонки."""
    if isinstance(value, date):
        return "00.00.0000"
    return "" if value is None else str(value)


def _format(cell: Any, kind: Kind, value: Any) -> None:
    """Выравнивание и формат числа по виду колонки."""
    if isinstance(value, date):
        cell.number_format = "DD.MM.YYYY"
        cell.alignment = Alignment(horizontal="center", vertical="top")
    elif kind == "percent":
        cell.value = value / 100 if isinstance(value, (int, float)) else _value(value)
        cell.number_format = "0.0%"
        cell.alignment = Alignment(horizontal="right", vertical="top")
    elif kind == "int":
        cell.alignment = Alignment(horizontal="right", vertical="top")
    elif kind == "center":
        cell.alignment = Alignment(horizontal="center", vertical="top")
    else:
        cell.alignment = Alignment(vertical="top", wrap_text=kind == "wrap")
