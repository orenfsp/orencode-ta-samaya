import re
from collections.abc import Sequence
from dataclasses import dataclass, field
from io import BytesIO
from typing import Literal
from uuid import UUID

from core.errors import UnsupportedFileType, ValidationFailed
from infrastructure.account.user import UserModel
from infrastructure.participation.contest_application import ContestApplicationModel
from results.domain.results import Results

MAX_ROWS = 2000
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DSQ = ("дискв", "dsq", "dq", "снят")
LINE = re.compile(r"^(?P<place>\d{1,3})[.)]?\s+(?P<name>\D.*?)(?:\s+(?P<points>\d+(?:[.,]\d+)?))?$")
DSQ_LINE = re.compile(r"^(?:дискв\S*|dsq|dq)\s+(?P<name>\D.*?)(?:\s+\d+(?:[.,]\d+)?)?$", re.IGNORECASE)
Status = Literal["matched", "doubtful", "unmatched"]


def normalize(text: str) -> str:
    """Готовит имя к сравнению: регистр, ё/е, знаки препинания и лишние пробелы не важны."""
    return " ".join(re.sub(r"[^0-9a-zа-я]+", " ", text.lower().replace("ё", "е")).split())


def person(user: UserModel) -> str:
    """ФИО человека для сравнения с протоколом."""
    return normalize(" ".join(part for part in (user.last_name, user.first_name, user.middle_name) if part))


@dataclass(slots=True)
class SourceRow:
    line: int
    name: str
    place: int | None
    disqualified: bool
    points: str | None


@dataclass(slots=True)
class ImportRow:
    line: int
    name: str
    place: int | None
    disqualified: bool
    points: str | None
    status: Status
    application_id: int | None
    candidates: list[int] = field(default_factory=list)
    note: str | None = None


@dataclass(frozen=True, slots=True)
class ImportPreview:
    file_name: str
    rows: list[ImportRow]
    entries: Sequence[ContestApplicationModel]
    missing: list[int]


class ProtocolImport:
    """Импорт протокола из xlsx или PDF: разбор строк и сопоставление с одобренными заявками по ФИО и команде."""

    def __init__(self, results: Results) -> None:
        """Сохраняет домен результатов: права и одобренные заявки берутся из него."""
        self._results = results

    async def preview(
        self, user_id: UUID, is_federation: bool, contest_id: int, file_name: str, content_type: str | None, data: bytes
    ) -> ImportPreview:
        """Разбирает файл и показывает совпадения, сомнительные и несопоставленные строки; ничего не сохраняет."""
        contest, entries, _ = await self._results.form(user_id, is_federation, contest_id)
        if contest.status not in Results.EDITABLE:
            raise ValidationFailed(fields={"file": ["Места вносятся, когда соревнование идёт или завершено"]})
        lowered = file_name.lower()
        if data[:4] == b"%PDF" or lowered.endswith(".pdf"):
            rows = self._pdf(data)
        elif data[:2] == b"PK" or lowered.endswith(".xlsx") or content_type == XLSX:
            rows = self._xlsx(data)
        else:
            raise UnsupportedFileType("Загрузите протокол в xlsx или PDF")
        if not rows:
            raise ValidationFailed(fields={"file": ["В файле не нашлось строк «место — участник»"]})
        if len(rows) > MAX_ROWS:
            raise ValidationFailed(fields={"file": [f"Не больше {MAX_ROWS} строк"]})
        matched = [self._match(row, entries) for row in rows]
        taken: dict[int, list[ImportRow]] = {}
        for row in matched:
            if row.application_id is not None:
                taken.setdefault(row.application_id, []).append(row)
        team: set[int] = set()
        for same in taken.values():
            if len(same) > 1 and len({(row.place, row.disqualified) for row in same}) == 1:
                team.update(id(row) for row in same[1:])
            elif len(same) > 1:
                for row in same:
                    row.status, row.note = "doubtful", "Одна и та же заявка в нескольких строках с разными местами"
        places = [row.place for row in matched if row.status == "matched" and row.place is not None and id(row) not in team]
        if len(places) != len(set(places)):
            for row in matched:
                if row.status == "matched" and row.place is not None and places.count(row.place) > 1:
                    row.status, row.note = "doubtful", "Место повторяется"
        missing = [entry.id for entry in entries if entry.id not in taken]
        return ImportPreview(file_name=file_name, rows=matched, entries=entries, missing=missing)

    @staticmethod
    def _match(row: SourceRow, entries: Sequence[ContestApplicationModel]) -> ImportRow:
        """Ищет заявку по названию команды и ФИО участников: полное совпадение, фамилия и имя, слова в другом порядке."""
        wanted = normalize(row.name)
        words = wanted.split()
        result = ImportRow(row.line, row.name, row.place, row.disqualified, row.points, "unmatched", None)
        exact, short, loose = [], [], []
        for entry in entries:
            members = [member.user for member in entry.members if member.is_active]
            names = [person(user) for user in members]
            if (entry.team_name and normalize(entry.team_name) == wanted) or wanted in names:
                exact.append(entry.id)
            elif any(name.split()[:2] == words[:2] for name in names if len(words) >= 2):
                short.append(entry.id)
            elif any(set(name.split()) >= set(words) and len(words) >= 2 for name in names):
                loose.append(entry.id)
        if len(exact) == 1:
            result.status, result.application_id = "matched", exact[0]
        elif exact or short or loose:
            candidates = exact or short or loose
            result.status, result.candidates = "doubtful", candidates
            result.application_id = candidates[0] if len(candidates) == 1 else None
            result.note = ("Несколько заявок с таким именем" if len(candidates) > 1
                           else "Совпали фамилия и имя" if short else "Совпали слова в другом порядке")
        else:
            result.note = "Нет одобренной заявки с таким участником или командой"
        return result

    @staticmethod
    def _xlsx(data: bytes) -> list[SourceRow]:
        """Читает первый лист: ищет заголовки «место», «ФИО / участник / команда», «баллы»; без них — две первые колонки."""
        from openpyxl import load_workbook

        try:
            book = load_workbook(BytesIO(data), read_only=True, data_only=True)
        except Exception as error:
            raise ValidationFailed(fields={"file": ["Файл не читается как xlsx"]}) from error
        sheet = book.worksheets[0]
        table = [["" if value is None else str(value).strip() for value in row] for row in sheet.iter_rows(values_only=True)]
        name_col, place_col, points_col, start = 1, 0, None, 0
        for index, cells in enumerate(table[:30]):
            heads = [normalize(cell) for cell in cells]
            place = next((i for i, head in enumerate(heads) if head.startswith("место")), None)
            name = next((i for i, head in enumerate(heads)
                         if any(word in head for word in ("фио", "участник", "спортсмен", "команда", "фамилия"))), None)
            if place is not None and name is not None:
                points = next((i for i, head in enumerate(heads) if head.startswith(("балл", "очк"))), None)
                name_col, place_col, points_col, start = name, place, points, index + 1
                break
        rows: list[SourceRow] = []
        for index, cells in enumerate(table[start:], start=start + 1):
            cells = cells + [""] * (max(name_col, place_col, points_col or 0) + 1 - len(cells))
            name, place = cells[name_col], cells[place_col].lower()
            if not name and not place:
                if rows:
                    break
                continue
            disqualified = place.startswith(DSQ)
            number = place.split(".")[0]
            if not name or not (number.isdigit() or disqualified):
                continue
            rows.append(SourceRow(index, name, None if disqualified else int(number), disqualified,
                                  cells[points_col] or None if points_col is not None else None))
        return rows

    @staticmethod
    def _pdf(data: bytes) -> list[SourceRow]:
        """Извлекает текст PDF и берёт строки вида «1. Иванов Иван Иванович 120», «Дисквалификация Петров Пётр»."""
        from pypdf import PdfReader

        try:
            lines = [line.strip() for page in PdfReader(BytesIO(data)).pages for line in (page.extract_text() or "").splitlines()]
        except Exception as error:
            raise ValidationFailed(fields={"file": ["Файл не читается как PDF"]}) from error
        rows: list[SourceRow] = []
        pending: int | None = None
        for index, line in enumerate(lines, start=1):
            if not line:
                continue
            if line.isdigit() and len(line) <= 3:
                pending = int(line)
                continue
            dsq = DSQ_LINE.match(line)
            found = LINE.match(line)
            if dsq:
                rows.append(SourceRow(index, dsq["name"].strip(), None, True, None))
            elif found and len(normalize(found["name"])) >= 3:
                rows.append(SourceRow(index, found["name"].strip(), int(found["place"]), False, found["points"]))
            elif pending is not None and re.match(r"^[А-ЯЁA-Z«\"]", line):
                tail = re.match(r"^(?P<name>.*?\D)(?:\s+(?P<points>\d+(?:[.,]\d+)?))?$", line)
                if tail and len(normalize(tail["name"])) >= 3:
                    rows.append(SourceRow(index, tail["name"].strip(), pending, False, tail["points"]))
            pending = None
        return rows
