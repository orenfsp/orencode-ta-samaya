from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

from infrastructure.contests.contest import ContestModel
from infrastructure.participation.contest_application import ContestApplicationModel
from infrastructure.participation.contest_result import ContestResultModel
from infrastructure.enums import PARTICIPATION_TYPE_TITLES, VERIFICATION_STATUS_TITLES, ParticipationType
from infrastructure.qualifications.judge_assignment import JudgeAssignmentModel
from infrastructure.account.user import UserModel


def _name(user: UserModel) -> str:
    """Собирает фамилию, имя и отчество пользователя в одну строку."""
    return " ".join(part for part in (user.last_name, user.first_name, user.middle_name) if part)


def _dates(starts: date, ends: date) -> str:
    """Форматирует одну дату или диапазон дат проведения соревнования."""
    return starts.strftime("%d.%m.%Y") if starts == ends else f"{starts:%d.%m.%Y} — {ends:%d.%m.%Y}"


@dataclass(frozen=True, slots=True)
class JudgeFile:
    judge: UserModel
    position: str
    kind: str
    name: str
    url: str
    comment: str
    removed: bool


class Protocol:
    """Формирует итоговый протокол соревнования в виде книги Excel."""

    HEADERS = ("Место", "Участник / команда", "Состав", "Баллы")
    WIDTHS = (10, 42, 60, 10)

    @classmethod
    def build(
        cls,
        contest: ContestModel,
        entries: Sequence[ContestApplicationModel],
        results: Sequence[ContestResultModel],
        judges: Sequence[JudgeAssignmentModel] = (),
        files: Sequence[JudgeFile] = (),
    ) -> bytes:
        """Строит xlsx-протокол соревнования: сведения, таблицу мест и судейскую коллегию."""
        book = Workbook()
        sheet = book.active
        assert sheet is not None
        sheet.title = "Протокол"
        sheet.append(["Протокол соревнования"])
        sheet["A1"].font = Font(bold=True, size=14)
        sheet.append([contest.title])
        sheet["A2"].font = Font(bold=True)
        place = ", ".join(part for part in (contest.municipality.name if contest.municipality else None, contest.venue) if part)
        facts = [
            ("Дисциплина", contest.discipline.name),
            ("Тип и уровень", f"{contest.contest_type.name}, {contest.contest_level.name}"),
            ("Формат", contest.contest_format.name),
            ("Даты проведения", _dates(contest.starts_on, contest.ends_on)),
            ("Место проведения", place or "—"),
            ("Участие", PARTICIPATION_TYPE_TITLES[contest.participation_type]),
            ("Организатор", _name(contest.organizer)),
            ("Верификация", VERIFICATION_STATUS_TITLES[contest.verification_status]),
        ]
        sheet.append([])
        for label, value in facts:
            sheet.append([label, value])
            sheet.cell(sheet.max_row, 1).font = Font(bold=True)
        sheet.append([])
        sheet.append(list(cls.HEADERS))
        header_row = sheet.max_row
        thin = Side(style="thin", color="C8C9CA")
        for column in range(1, len(cls.HEADERS) + 1):
            cell = sheet.cell(header_row, column)
            cell.font = Font(bold=True)
            cell.border = Border(bottom=thin)
        team = contest.participation_type == ParticipationType.TEAM
        by_application: dict[int, ContestResultModel] = {}
        for result in results:
            by_application.setdefault(result.application_id, result)
        rows = []
        for entry in entries:
            result = by_application.get(entry.id)
            members = [member.user for member in entry.members if member.is_active]
            title = entry.team_name or (_name(members[0]) if members else "—")
            roster = ", ".join(_name(user) for user in members) if team else ""
            if result is None:
                rows.append((2, 0, ["", title, roster, ""]))
            elif result.place is None:
                rows.append((1, 0, ["Дисквалификация", title, roster, 0]))
            else:
                rows.append((0, result.place, [result.place, title, roster, result.points]))
        for _, _, values in sorted(rows, key=lambda row: (row[0], row[1])):
            sheet.append(values)
            for column in (1, 4):
                sheet.cell(sheet.max_row, column).alignment = Alignment(horizontal="center")
            sheet.cell(sheet.max_row, 3).alignment = Alignment(wrap_text=True, vertical="top")
        if judges:
            sheet.append([])
            sheet.append(["Судейская коллегия"])
            sheet.cell(sheet.max_row, 1).font = Font(bold=True)
            for assignment in judges:
                sheet.append(["", assignment.position.name, _name(assignment.judge)])
        if files:
            sheet.append([])
            sheet.append(["Документы судейской коллегии"])
            sheet.cell(sheet.max_row, 1).font = Font(bold=True)
            for item in files:
                judge = _name(item.judge) + (" (снят)" if item.removed else "")
                note = " · ".join(part for part in (item.position, item.comment) if part)
                sheet.append([item.kind, judge, note, f'=HYPERLINK("{item.url}", "{item.name.replace(chr(34), "")}")'])
                sheet.cell(sheet.max_row, 4).font = Font(color="402FFF", underline="single")
        for index, width in enumerate(cls.WIDTHS, start=1):
            sheet.column_dimensions[chr(64 + index)].width = width
        buffer = BytesIO()
        book.save(buffer)
        return buffer.getvalue()
