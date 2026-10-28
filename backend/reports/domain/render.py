from io import BytesIO
from pathlib import Path

from jinja2 import Environment, FileSystemLoader, select_autoescape
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from reports.domain.layout import facts, sections
from reports.domain.person_report import PersonReport, full_name

ROOT = Path(__file__).resolve().parent.parent
ASSETS = ROOT.parent / "documents" / "assets"
TEMPLATES = Environment(loader=FileSystemLoader(ROOT / "templates"), autoescape=select_autoescape(["html"]))
THIN = Side(style="thin", color="C8C9CA")
HEAD = PatternFill("solid", fgColor="EDEDED")


def summary(report: PersonReport) -> list[tuple[int, str]]:
    """Считает сводные числа отчёта: соревнования, призовые места, баллы, документы."""
    places = [item.place for item in report.starts if item.place is not None]
    return [
        (len(report.starts), "соревнований"),
        (sum(1 for place in places if place <= 3), "призовых мест"),
        (report.points, "баллов за период"),
        (len(report.documents), "документов"),
    ]


def pdf(report: PersonReport) -> bytes:
    """Рендерит отчёт по человеку в PDF по HTML-шаблону."""
    html = TEMPLATES.get_template("person_report.html").render(
        name=full_name(report.user), facts=facts(report), summary=summary(report), sections=sections(report)
    )
    from weasyprint import HTML

    return HTML(string=html, base_url=str(ASSETS) + "/").write_pdf() or b""


def xlsx(report: PersonReport) -> bytes:
    """Рендерит отчёт по человеку в книгу xlsx: сводный лист и лист на каждый раздел."""
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Сводка"
    sheet.cell(1, 1, "Отчёт по человеку").font = Font(bold=True, size=14)
    row = 3
    for label, value in facts(report):
        sheet.cell(row, 1, label).font = Font(color="5B5C62")
        sheet.cell(row, 2, value)
        row += 1
    row += 1
    for value, caption in summary(report):
        sheet.cell(row, 1, caption).font = Font(color="5B5C62")
        sheet.cell(row, 2, value).font = Font(bold=True)
        row += 1
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 48
    for section in sections(report):
        page = book.create_sheet(section.title[:31])
        for column, header in enumerate(section.headers, start=1):
            cell = page.cell(1, column, header)
            cell.font = Font(bold=True)
            cell.fill = HEAD
            cell.border = Border(bottom=THIN)
        if not section.rows:
            page.cell(2, 1, section.empty).font = Font(color="5B5C62")
        widths = [len(header) for header in section.headers]
        for index, values in enumerate(section.rows, start=2):
            for column, value in enumerate(values, start=1):
                cell = page.cell(index, column, value)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                widths[column - 1] = max(widths[column - 1], min(len(str(value)), 60))
        for column, width in enumerate(widths, start=1):
            page.column_dimensions[get_column_letter(column)].width = width + 3
        page.freeze_panes = "A2"
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
