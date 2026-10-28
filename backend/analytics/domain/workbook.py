from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from analytics.domain.analytics import TITLES, Query, Result

THIN = Side(style="thin", color="BFC3CC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = PatternFill("solid", fgColor="E6E9F0")


def filename(query: Query, block: str | None = None) -> str:
    """Формирует имя xlsx-файла по периоду запроса и, если указан, блоку."""
    period = query.period
    name = f"analytics-{block}" if block else "analytics"
    return f"{name}-{period.start:%Y-%m-%d}-{period.end:%Y-%m-%d}.xlsx"


def _sheet_title(title: str) -> str:
    # Excel не пускает в имя листа []:*?/\ и длиннее 31 знака
    """Приводит заголовок блока к допустимому имени листа Excel."""
    return "".join(char for char in title if char not in "[]:*?/\\")[:31]


def _header(sheet: Worksheet, title: str, parameters: list[list[Any]]) -> int:
    """Пишет заголовок и строки параметров на лист, возвращает номер следующей свободной строки."""
    sheet.cell(1, 1, title).font = Font(bold=True, size=14)
    row = 2
    for label, value in parameters:
        sheet.cell(row, 1, label).font = Font(color="5A6070")
        sheet.cell(row, 2, value)
        row += 1
    return row + 1


def _fill(sheet: Worksheet, result: Result, start: int) -> None:
    """Выводит все таблицы результата на лист начиная со строки start, подбирая ширины колонок."""
    widths: dict[int, int] = {}
    row = start
    for table in result.tables:
        sheet.cell(row, 1, table.title).font = Font(bold=True, size=12)
        row += 1
        for column, header in enumerate(table.headers, 1):
            cell = sheet.cell(row, column, header)
            cell.font = Font(bold=True)
            cell.fill = HEAD
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            widths[column] = max(widths.get(column, 0), len(header))
        row += 1
        for values in table.rows:
            for column, value in enumerate(values, 1):
                index = column - 1
                if index in table.percent_columns and isinstance(value, (int, float)):
                    cell = sheet.cell(row, column, value / 100)
                    cell.number_format = "0.0%"
                else:
                    cell = sheet.cell(row, column, "" if value is None else value)
                cell.border = BORDER
                widths[column] = max(widths.get(column, 0), len(str(value)))
            row += 1
        if not table.rows:
            sheet.cell(row, 1, "Нет данных").font = Font(italic=True, color="5A6070")
            row += 1
        row += 1
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = min(max(width + 2, 12), 60)


def build(results: list[Result], parameters: list[list[Any]], with_parameters_sheet: bool) -> bytes:
    """Собирает книгу xlsx: по листу на каждый результат и опционально лист параметров; возвращает байты файла."""
    book = Workbook()
    first = book.active
    assert first is not None
    if with_parameters_sheet:
        first.title = "Параметры"
        _header(first, "Аналитика ФСП Оренбургской области", parameters)
        first.column_dimensions["A"].width = 18
        first.column_dimensions["B"].width = 40
    for index, result in enumerate(results):
        title = TITLES[result.block]
        sheet = first if index == 0 and not with_parameters_sheet else book.create_sheet()
        sheet.title = _sheet_title(title)
        _fill(sheet, result, _header(sheet, title, parameters))
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
